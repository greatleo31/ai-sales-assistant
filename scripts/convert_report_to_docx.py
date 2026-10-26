import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_run_font(run, font_name="宋体", ascii_font="Times New Roman", size_pt=12, bold=False, italic=False, color=None):
    run.font.name = ascii_font
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), ascii_font)
    rFonts.set(qn('w:hAnsi'), ascii_font)
    rFonts.set(qn('w:eastAsia'), font_name)
    rPr.append(rFonts)

def set_cell_background(cell, fill_hex="F2F4F7"):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_three_line_borders(table):
    """设置标准学术三线表边框：顶线1.5pt，底线1.5pt，表头线0.75pt，无竖线"""
    tblPr = table._tbl.tblPr
    # 移除现有边框
    for child in list(tblPr):
        if child.tag.endswith('tblBorders'):
            tblPr.remove(child)
    
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>\n'
        f'  <w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:insideH w:val="none"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

    # 给表头行单独加底部细横线
    if len(table.rows) > 0:
        for cell in table.rows[0].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>\n'
                f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
                f'</w:tcBorders>'
            )
            tcPr.append(tcBorders)

def set_grid_borders(table):
    """设置元数据全实线网格边框"""
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag.endswith('tblBorders'):
            tblPr.remove(child)
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="D0D5DD"/>\n'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="D0D5DD"/>\n'
        f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="D0D5DD"/>\n'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="D0D5DD"/>\n'
        f'  <w:insideH w:val="single" w:sz="6" w:space="0" w:color="D0D5DD"/>\n'
        f'  <w:insideV w:val="single" w:sz="6" w:space="0" w:color="D0D5DD"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)


def parse_inline_markdown(para, text, default_font="宋体", default_ascii="Times New Roman", default_size=12, default_color=None):
    # 处理 **bold**, `code`, *italic*
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**') and len(tok) >= 4:
            run = para.add_run(tok[2:-2])
            set_run_font(run, font_name="黑体" if default_font=="宋体" else default_font, ascii_font=default_ascii, size_pt=default_size, bold=True, color=default_color)
        elif tok.startswith('`') and tok.endswith('`') and len(tok) >= 2:
            run = para.add_run(tok[1:-1])
            set_run_font(run, font_name="Consolas", ascii_font="Consolas", size_pt=default_size*0.9, color=RGBColor(0x0F, 0x52, 0xBA))
        elif tok.startswith('*') and tok.endswith('*') and len(tok) >= 2:
            run = para.add_run(tok[1:-1])
            set_run_font(run, font_name=default_font, ascii_font=default_ascii, size_pt=default_size, italic=True, color=default_color)
        else:
            run = para.add_run(tok)
            set_run_font(run, font_name=default_font, ascii_font=default_ascii, size_pt=default_size, color=default_color)


def convert_md_to_docx(md_path, docx_path, img_base_dir):
    doc = Document()
    
    # 页面设置：A4纸，标准1英寸边距
    sections = doc.sections
    for section in sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table(lines_to_flush):
        if not lines_to_flush:
            return
        # 过滤分隔行 | :--- | :--- |
        cleaned_rows = []
        for line in lines_to_flush:
            raw = line.strip().strip('|')
            parts = [c.strip() for c in raw.split('|')]
            if any('---' in c for c in parts):
                continue
            cleaned_rows.append(parts)
        
        if not cleaned_rows:
            return

        is_meta_table = any("学生姓名" in "".join(r) or "任课教师" in "".join(r) for r in cleaned_rows)
        
        col_count = max(len(r) for r in cleaned_rows)
        table = doc.add_table(rows=len(cleaned_rows), cols=col_count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        for r_idx, row_data in enumerate(cleaned_rows):
            for c_idx in range(col_count):
                cell = table.cell(r_idx, c_idx)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                val = row_data[c_idx] if c_idx < len(row_data) else ""
                
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                
                # 样式控制
                if is_meta_table:
                    if r_idx % 2 == 0 or c_idx % 2 == 0:
                        set_cell_background(cell, "F8F9FA")
                    parse_inline_markdown(p, val, default_font="宋体", default_size=10.5)
                else:
                    if r_idx == 0:
                        set_cell_background(cell, "F2F4F7")
                        parse_inline_markdown(p, val, default_font="黑体", default_size=10.5, default_color=RGBColor(0x10, 0x18, 0x28))
                    else:
                        parse_inline_markdown(p, val, default_font="宋体", default_size=10)

        if is_meta_table:
            set_grid_borders(table)
        else:
            set_three_line_borders(table)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def flush_code(lines_to_flush):
        if not lines_to_flush:
            return
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = True
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F7F8FA")
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        
        # 细灰框
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="EAECF0"/>\n'
            f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="EAECF0"/>\n'
            f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="EAECF0"/>\n'
            f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="EAECF0"/>\n'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        
        full_text = "".join(lines_to_flush)
        run = p.add_run(full_text)
        set_run_font(run, font_name="Consolas", ascii_font="Consolas", size_pt=9.5, color=RGBColor(0x34, 0x40, 0x54))
        
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 处理代码块
        if stripped.startswith("```"):
            if in_code_block:
                in_code_block = False
                flush_code(code_lines)
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # 处理表格
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_lines.append(line)
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                flush_table(table_lines)
                table_lines = []

        # 空行或分隔线
        if not stripped or stripped == "---":
            i += 1
            continue

        # 匹配大标题 #
        if stripped.startswith("# ") and not stripped.startswith("## "):
            title_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(10)
            parse_inline_markdown(p, title_text, default_font="黑体", default_size=20, default_color=RGBColor(0x10, 0x18, 0x28))
            i += 1
            continue

        # 匹配一级标题 ##
        if stripped.startswith("## ") and not stripped.startswith("### "):
            h1_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            parse_inline_markdown(p, h1_text, default_font="黑体", default_size=15, default_color=RGBColor(0x10, 0x18, 0x28))
            i += 1
            continue

        # 匹配二级标题 ###
        if stripped.startswith("### "):
            h2_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            parse_inline_markdown(p, h2_text, default_font="黑体", default_size=13, default_color=RGBColor(0x1D, 0x29, 0x39))
            i += 1
            continue

        # 匹配三级标题 ####
        if stripped.startswith("#### "):
            h3_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            parse_inline_markdown(p, h3_text, default_font="黑体", default_size=12, default_color=RGBColor(0x34, 0x40, 0x54))
            i += 1
            continue

        # 处理图片 ![alt](path)
        img_match = re.match(r'!\[(.*?)\]\((.*?)\)', stripped)
        if img_match:
            alt_text = img_match.group(1)
            img_rel_path = img_match.group(2)
            
            # 解析图片物理路径
            if img_rel_path.startswith("../eval_reports/"):
                actual_img_path = os.path.join(img_base_dir, "eval_reports", os.path.basename(img_rel_path))
            else:
                actual_img_path = os.path.abspath(os.path.join(os.path.dirname(md_path), img_rel_path))

            if os.path.exists(actual_img_path):
                img_para = doc.add_paragraph()
                img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_para.paragraph_format.space_before = Pt(8)
                img_para.paragraph_format.space_after = Pt(2)
                run = img_para.add_run()
                # 宽度适配 5.8 英寸
                run.add_picture(actual_img_path, width=Inches(5.8))
                
                # 检查下一行是否是图题
                if i + 1 < len(lines) and (lines[i+1].strip().startswith("图") or lines[i+1].strip().startswith("*图")):
                    cap_text = lines[i+1].strip().strip('*')
                    cap_para = doc.add_paragraph()
                    cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap_para.paragraph_format.space_before = Pt(2)
                    cap_para.paragraph_format.space_after = Pt(8)
                    cap_para.paragraph_format.keep_with_next = False
                    run_cap = cap_para.add_run(cap_text)
                    set_run_font(run_cap, font_name="宋体", ascii_font="Times New Roman", size_pt=10.5, bold=True, color=RGBColor(0x47, 0x54, 0x67))
                    i += 2
                    continue
            i += 1
            continue

        # 检查是否为独立图题行或表题行
        if stripped.startswith("表 ") or stripped.startswith("表1 ") or stripped.startswith("表2 ") or stripped.startswith("表3 ") or stripped.startswith("表4 "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped)
            set_run_font(run, font_name="黑体", ascii_font="Times New Roman", size_pt=10.5, bold=True)
            i += 1
            continue

        # 检查独立公式行 $$ ... $$
        if stripped.startswith("$$") and stripped.endswith("$$") and len(stripped) >= 4:
            formula_text = stripped[2:-2].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(formula_text)
            set_run_font(run, font_name="Times New Roman", ascii_font="Times New Roman", size_pt=11, italic=True)
            i += 1
            continue

        # 普通正文段落
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        
        # 列表或普通缩进
        if stripped.startswith("- ") or re.match(r'^\d+\.\s', stripped):
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(0)
        elif not stripped.startswith("图") and not stripped.startswith("表") and not stripped.startswith("摘要") and not stripped.startswith("【"):
            p.paragraph_format.first_line_indent = Inches(0.28) # 缩进约2汉字

        parse_inline_markdown(p, stripped, default_font="宋体", default_size=11, default_color=RGBColor(0x1D, 0x29, 0x39))
        i += 1

    if in_table:
        flush_table(table_lines)
    if in_code_block:
        flush_code(code_lines)

    doc.save(docx_path)
    print(f"Saved docx to: {docx_path}")

if __name__ == "__main__":
    md_file = r"D:\projectCollection\dw\ai-sales-assistant\docs\reports\huigou_reasoning_research_report.md"
    docx_file = r"D:\projectCollection\dw\ai-sales-assistant\docs\reports\基于神经-符号双轨推理的惠购智能导购与售后争议治理系统_课程报告.docx"
    root_docs = r"D:\projectCollection\dw\ai-sales-assistant\docs"
    convert_md_to_docx(md_file, docx_file, root_docs)
